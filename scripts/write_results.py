import subprocess
import re
import openpyxl
from datetime import datetime
from openpyxl.styles import PatternFill, Font, Alignment

inventory = "/home/controller/linux-server-operations-lab/inventory/hosts_generated.ini"
excel_path = "/media/sf_Ansible_Python/Linux_Server_Operations_Lab.xlsx"


# ========================================
# Disk使用率を取得
# ========================================
disk_result = subprocess.run(
    [
        "ansible",
        "servers",
        "-i",
        inventory,
        "-a",
        "df -P /",
        "-T",
        "5"
    ],
    capture_output=True,
    text=True,
    timeout=30
)

disk_status = {}
current_host = None

for line in disk_result.stdout.splitlines():

    host_match = re.match(r"^(\S+) \|", line)

    if host_match:
        current_host = host_match.group(1)

        if "UNREACHABLE" in line:
            disk_status[current_host] = {
                "percent": None,
                "status": "UNREACHABLE"
            }

        elif "FAILED" in line:
            disk_status[current_host] = {
                "percent": None,
                "status": "FAILED"
            }

    disk_match = re.search(r"(\d+)%\s+/$", line)

    if disk_match and current_host:
        disk_percent = int(disk_match.group(1))

        if disk_percent >= 80:
            status = "WARNING"
        else:
            status = "OK"

        disk_status[current_host] = {
            "percent": disk_percent,
            "status": status
        }


# ========================================
# nginx状態を取得
# ========================================
nginx_result = subprocess.run(
    [
        "ansible",
        "servers",
        "-i",
        inventory,
        "-a",
        "systemctl is-active nginx",
        "-T",
        "5"
    ],
    capture_output=True,
    text=True,
    timeout=30
)

nginx_status = {}
current_host = None

for line in nginx_result.stdout.splitlines():

    host_match = re.match(r"^(\S+) \|", line)

    if host_match:
        current_host = host_match.group(1)

        if "UNREACHABLE" in line:
            nginx_status[current_host] = "UNREACHABLE"

        elif "FAILED" in line:
            nginx_status[current_host] = "FAILED"

    if line.strip() == "active" and current_host:
        nginx_status[current_host] = "OK"


# ========================================
# Excelを開く
# ========================================
wb = openpyxl.load_workbook(excel_path)

master_ws = wb["01_SERVER_MASTER"]
result_ws = wb["02_CHECK_RESULTS"]


# ========================================
# Hostname → ServerID
# ========================================
server_ids = {}

for row in master_ws.iter_rows(min_row=2, values_only=True):

    server_id = row[0]
    hostname = row[1]

    if hostname:
        server_ids[hostname] = server_id


# ========================================
# 古い結果を削除
# ========================================
for row in result_ws.iter_rows(
    min_row=2,
    max_row=result_ws.max_row,
    min_col=1,
    max_col=6
):
    for cell in row:
        cell.value = None
        cell.fill = PatternFill(fill_type=None)


# ========================================
# 表示用スタイル
# ========================================
ok_fill = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

warning_fill = PatternFill(
    fill_type="solid",
    fgColor="FFEB9C"
)

error_fill = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)

header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

bold_font = Font(bold=True)


# ========================================
# ヘッダーを整える
# ========================================
for cell in result_ws[1]:
    cell.font = bold_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")


# ========================================
# 最終結果を作成
# ========================================
hosts = set(disk_status.keys()) | set(nginx_status.keys())

# ServerID順に並べる
hosts = sorted(
    hosts,
    key=lambda host: server_ids.get(host, "ZZZ")
)

row_number = 2
executed_at = datetime.now()

for host in hosts:

    disk = disk_status.get(
        host,
        {
            "percent": None,
            "status": "FAILED"
        }
    )

    nginx = nginx_status.get(
        host,
        "FAILED"
    )

    if disk["status"] == "UNREACHABLE" or nginx == "UNREACHABLE":
        final = "UNREACHABLE"

    elif disk["status"] == "FAILED" or nginx == "FAILED":
        final = "FAILED"

    elif disk["status"] == "WARNING":
        final = "WARNING"

    else:
        final = "OK"

    server_id = server_ids.get(host, "UNKNOWN")

    if disk["percent"] is None:
        disk_display = "-"
    else:
        disk_display = f'{disk["percent"]}%'

    result_ws.cell(row=row_number, column=1, value=server_id)
    result_ws.cell(row=row_number, column=2, value=host)
    result_ws.cell(row=row_number, column=3, value=disk_display)
    result_ws.cell(row=row_number, column=4, value=nginx)
    result_ws.cell(row=row_number, column=5, value=final)
    result_ws.cell(row=row_number, column=6, value=executed_at)

    # 日時表示
    result_ws.cell(
        row=row_number,
        column=6
    ).number_format = "yyyy-mm-dd hh:mm:ss"

    # Resultを色分け
    result_cell = result_ws.cell(
        row=row_number,
        column=5
    )

    if final == "OK":
        result_cell.fill = ok_fill

    elif final == "WARNING":
        result_cell.fill = warning_fill

    else:
        result_cell.fill = error_fill

    print(
        f"{server_id} "
        f"{host} "
        f"Disk={disk_display} "
        f"nginx={nginx} "
        f"Result={final}"
    )

    row_number += 1


# ========================================
# 列幅を調整
# ========================================
result_ws.column_dimensions["A"].width = 12
result_ws.column_dimensions["B"].width = 15
result_ws.column_dimensions["C"].width = 10
result_ws.column_dimensions["D"].width = 10
result_ws.column_dimensions["E"].width = 15
result_ws.column_dimensions["F"].width = 22


# ========================================
# Excel保存
# ========================================
wb.save(excel_path)

print()
print("Excel write completed")
