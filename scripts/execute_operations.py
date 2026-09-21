import subprocess
import openpyxl
from datetime import datetime
from openpyxl.styles import PatternFill

excel_path = "/media/sf_Ansible_Python/Linux_Server_Operations_Lab.xlsx"
inventory = "/home/controller/linux-server-operations-lab/inventory/hosts_generated.ini"


# ========================================
# Excelを開く
# ========================================
wb = openpyxl.load_workbook(excel_path)

master_ws = wb["01_SERVER_MASTER"]
operation_ws = wb["03_OPERATION_REQUESTS"]


# ========================================
# Status表示用の色
# ========================================
success_fill = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

failed_fill = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)


# ========================================
# ServerID → Hostname の対応表を作る
# ========================================
server_map = {}

for row in master_ws.iter_rows(min_row=2, values_only=True):

    server_id = row[0]
    hostname = row[1]

    if server_id and hostname:
        server_map[server_id] = hostname


# ========================================
# 操作指示を読み込む
# ========================================
for row_number in range(2, operation_ws.max_row + 1):

    request_id = operation_ws.cell(
        row=row_number,
        column=1
    ).value

    server_id = operation_ws.cell(
        row=row_number,
        column=2
    ).value

    action = operation_ws.cell(
        row=row_number,
        column=3
    ).value

    execute = operation_ws.cell(
        row=row_number,
        column=4
    ).value

    status = operation_ws.cell(
        row=row_number,
        column=5
    ).value


    # ========================================
    # YES以外は実行しない
    # ========================================
    if execute != "YES":
        continue


    # ========================================
    # SUCCESS済みは再実行しない
    # ========================================
    if status == "SUCCESS":
        continue


    # ========================================
    # ServerIDからHostnameを取得
    # ========================================
    hostname = server_map.get(server_id)

    if hostname is None:

        operation_ws.cell(
            row=row_number,
            column=5,
            value="FAILED"
        )

        operation_ws.cell(
            row=row_number,
            column=5
        ).fill = failed_fill

        operation_ws.cell(
            row=row_number,
            column=6,
            value="ServerID not found"
        )

        operation_ws.cell(
            row=row_number,
            column=7,
            value=datetime.now()
        )

        operation_ws.cell(
            row=row_number,
            column=7
        ).number_format = "yyyy-mm-dd hh:mm:ss"

        continue


    # ========================================
    # Action判定
    # ========================================
    if action == "CREATE_TEST_FILE":

        command = [
            "ansible",
            hostname,
            "-i",
            inventory,
            "-m",
            "file",
            "-a",
            "path=/tmp/phase8_test.txt state=touch"
        ]

    elif action == "DELETE_TEST_FILE":

        command = [
            "ansible",
            hostname,
            "-i",
            inventory,
            "-m",
            "file",
            "-a",
            "path=/tmp/phase8_test.txt state=absent"
        ]

    elif action == "RESTART_NGINX":

        command = [
            "ansible",
            hostname,
            "-i",
            inventory,
            "-a",
            "sudo -n /usr/bin/systemctl restart nginx"
        ]

    else:

        operation_ws.cell(
            row=row_number,
            column=5,
            value="FAILED"
        )

        operation_ws.cell(
            row=row_number,
            column=5
        ).fill = failed_fill

        operation_ws.cell(
            row=row_number,
            column=6,
            value="Unknown action"
        )

        operation_ws.cell(
            row=row_number,
            column=7,
            value=datetime.now()
        )

        operation_ws.cell(
            row=row_number,
            column=7
        ).number_format = "yyyy-mm-dd hh:mm:ss"

        continue


    # ========================================
    # Ansible実行
    # ========================================
    print(
        f"Executing {request_id}: "
        f"{server_id} -> {hostname} -> {action}"
    )

    result = subprocess.run(
        command,
        capture_output=True,
        text=True
    )


    # ========================================
    # 実行結果判定
    # ========================================
    if result.returncode == 0:

        operation_ws.cell(
            row=row_number,
            column=5,
            value="SUCCESS"
        )

        operation_ws.cell(
            row=row_number,
            column=5
        ).fill = success_fill

        operation_ws.cell(
            row=row_number,
            column=6,
            value=f"{action} completed"
        )

        print("SUCCESS")

    else:

        operation_ws.cell(
            row=row_number,
            column=5,
            value="FAILED"
        )

        operation_ws.cell(
            row=row_number,
            column=5
        ).fill = failed_fill

        message = result.stderr.strip()

        if not message:
            message = result.stdout.strip()

        operation_ws.cell(
            row=row_number,
            column=6,
            value=message[:200]
        )

        print("FAILED")


    # ========================================
    # 実行日時
    # ========================================
    operation_ws.cell(
        row=row_number,
        column=7,
        value=datetime.now()
    )

    operation_ws.cell(
        row=row_number,
        column=7
    ).number_format = "yyyy-mm-dd hh:mm:ss"


# ========================================
# Excel保存
# ========================================
wb.save(excel_path)

print()
print("Operation processing completed")
